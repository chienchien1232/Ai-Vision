"""Generate the project's Word plan from the reviewable Markdown source.

Run using the bundled workspace Python with python-docx; no runtime dependency
is added to the Android/backend product.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/ke-hoach-bon-chuc-nang-2026-09-28.md"
OUTPUT = ROOT / "docs/Ke_hoach_hoan_thien_4_chuc_nang_kinh_AI_28-09-2026.docx"


def add_table(doc, lines):
    rows = [[part.strip() for part in line.strip().strip("|").split("|")]
            for line in lines if not re.match(r"^\|\s*[-:]", line)]
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.7, 1.45, 2.2, 1.55]
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    props.append(borders)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        tr_pr = table.rows[-1]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for column, (cell, text) in enumerate(zip(cells, row)):
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCE6F1" if index == 0 else "F4F6F8" if index % 2 == 0 else "FFFFFF")
            tc_pr.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tc_pr.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(text)
            run.font.size = Pt(10.5)
            run.bold = index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == "Normal" else 22 if name == "Title" else 15 if name == "Heading 1" else 12)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.12
        for borders in style.element.findall(".//" + qn("w:pBdr")):
            borders.getparent().remove(borders)
    doc.styles["Normal"].paragraph_format.widow_control = True
    doc.core_properties.title = "Kế hoạch hoàn thiện bốn chức năng kính AI"
    doc.core_properties.author = "Ai Vision"
    doc.core_properties.subject = "Wake giọng nói local OCR và hỏi cảnh"
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    table_lines = []
    for line in [*lines, ""]:
        if line.startswith("|"):
            table_lines.append(line)
            continue
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []
        if not line:
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], "List Bullet")
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", line), "List Number")
        else:
            doc.add_paragraph(line)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
